import io
from pathlib import Path

import docx
import pymupdf
import pytest
from ebooklib import epub
from PIL import Image

from pdf2video.parsers import ParseError, load_document


def _paragraph(n: int) -> str:
    return " ".join(f"Sentence number {i} of the text is here." for i in range(n))


def test_markdown_chapters(tmp_path: Path):
    f = tmp_path / "book.md"
    f.write_text(
        "# My Book\n\nIntro text that is a preface.\n\n## First\n\nHello **world** "
        "and [a link](http://x).\n\n### Sub\n\nMore.\n\n## Second\n\n- item one\n- item two\n"
    )
    doc = load_document(f)
    assert doc.title == "My Book"
    assert [c.title for c in doc.chapters] == ["My Book", "First", "Second"]
    assert doc.chapters[1].paragraphs[0] == "Hello world and a link."
    assert "Sub" in doc.chapters[1].paragraphs


def test_plain_text_chapter_lines(tmp_path: Path):
    f = tmp_path / "novel.txt"
    f.write_text("CHAPTER I\n\nIt was a dark night.\n\nCHAPTER II\n\nThe sun rose.\n")
    doc = load_document(f)
    assert [c.title for c in doc.chapters] == ["CHAPTER I", "CHAPTER II"]


def test_docx_headings(tmp_path: Path):
    d = docx.Document()
    d.core_properties.title = "Doc Title"
    d.add_heading("Alpha", level=1)
    d.add_paragraph(_paragraph(3))
    d.add_heading("Alpha detail", level=2)
    d.add_paragraph("Detail text.")
    d.add_heading("Beta", level=1)
    d.add_paragraph("Beta text.")
    f = tmp_path / "test.docx"
    d.save(f)
    doc = load_document(f)
    assert doc.title == "Doc Title"
    assert [c.title for c in doc.chapters] == ["Alpha", "Beta"]
    assert "Alpha detail" in doc.chapters[0].paragraphs


def test_epub_toc(tmp_path: Path):
    book = epub.EpubBook()
    book.set_identifier("id1")
    book.set_title("Epub Title")
    book.add_author("Someone")
    png = io.BytesIO()
    Image.new("RGB", (300, 200), "red").save(png, "PNG")
    book.add_item(
        epub.EpubImage(file_name="img/pic.png", media_type="image/png", content=png.getvalue())
    )
    chapters = []
    figures = [
        '<figure><img src="img/pic.png" alt="ignored"/><figcaption>Fig. 1 A red flag</figcaption>'
        "</figure>",
        '<div><img src="img/pic.png" alt="A red square"/></div>',
    ]
    for i, (name, figure) in enumerate(zip(["One", "Two"], figures, strict=True), 1):
        c = epub.EpubHtml(title=name, file_name=f"c{i}.xhtml")
        c.content = f"<html><body><h1>{name}</h1><p>{_paragraph(2)}</p>{figure}</body></html>"
        book.add_item(c)
        chapters.append(c)
    book.toc = chapters
    book.spine = ["nav", *chapters]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    f = tmp_path / "test.epub"
    epub.write_epub(str(f), book)
    doc = load_document(f)
    assert doc.title == "Epub Title"
    assert doc.author == "Someone"
    assert [c.title for c in doc.chapters] == ["One", "Two"]
    assert doc.chapters[0].paragraphs[0].startswith("Sentence number 0")
    assert len(doc.chapters[1].images) == 1
    assert doc.chapters[0].images[0].caption == "Fig. 1 A red flag"
    assert doc.chapters[1].images[0].caption == "A red square"


def _png(size=(300, 200), color="red") -> bytes:
    data = io.BytesIO()
    Image.new("RGB", size, color).save(data, "PNG")
    return data.getvalue()


def test_docx_image_alt_text(tmp_path: Path):
    d = docx.Document()
    d.add_heading("Alpha", level=1)
    d.add_paragraph(_paragraph(3))
    d.add_picture(io.BytesIO(_png()))
    d.inline_shapes[0]._inline.docPr.set("descr", "A red square")
    d.add_heading("Beta", level=1)
    d.add_paragraph("Beta text.")
    f = tmp_path / "pics.docx"
    d.save(f)
    doc = load_document(f)
    alpha = next(c for c in doc.chapters if c.title == "Alpha")
    assert [i.caption for i in alpha.images] == ["A red square"]
    assert Image.open(io.BytesIO(alpha.images[0].data)).size == (300, 200)


def test_pdf_image_caption(tmp_path: Path):
    pdf = pymupdf.open()
    for title in ["First", "Second"]:
        page = pdf.new_page()
        page.insert_text((72, 90), title, fontsize=24)
        page.insert_text((72, 120), f"Some body text of {title} with words.", fontsize=11)
    page = pdf[1]
    page.insert_image(pymupdf.Rect(72, 150, 372, 350), stream=_png())
    page.insert_text((72, 370), "Figure 2: A red rectangle", fontsize=10)
    pdf.set_toc([[1, "First", 1], [1, "Second", 2]])
    f = tmp_path / "figs.pdf"
    pdf.save(f)
    doc = load_document(f)
    assert doc.chapters[0].images == []
    assert [i.caption for i in doc.chapters[1].images] == ["Figure 2: A red rectangle"]


def _make_pdf(path: Path, with_toc: bool) -> None:
    pdf = pymupdf.open()
    toc = []
    for i, title in enumerate(["Introduction", "Methods", "Results"], 1):
        page = pdf.new_page()
        page.insert_text((72, 90), title, fontsize=24)
        y = 130
        for line in range(12):
            page.insert_text((72, y), f"Body line {line} of {title} with some words.", fontsize=11)
            y += 16
        toc.append([1, title, i])
    if with_toc:
        pdf.set_toc(toc)
    pdf.set_metadata({"title": "PDF Title", "author": "A. Author"})
    pdf.save(path)


@pytest.mark.parametrize("with_toc", [True, False])
def test_pdf_chapters(tmp_path: Path, with_toc: bool):
    f = tmp_path / "paper.pdf"
    _make_pdf(f, with_toc)
    doc = load_document(f)
    assert doc.title == "PDF Title"
    assert [c.title for c in doc.chapters] == ["Introduction", "Methods", "Results"]
    assert "Body line 0 of Methods" in doc.chapters[1].text
    assert "Methods" not in doc.chapters[0].text


def test_unsupported(tmp_path: Path):
    f = tmp_path / "x.xyz"
    f.write_text("hi")
    with pytest.raises(ParseError):
        load_document(f)


def test_malformed_pdf_is_repaired_quietly(tmp_path: Path, capfd, caplog):
    import logging
    import re

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 90), "Chapter One", fontsize=24)
    page.insert_text((72, 130), "Readable body text survives the repair.", fontsize=11)
    raw = pdf.tobytes(deflate=False)
    # An invalid (non-name) key in the font dictionary, like buggy PDF generators write.
    raw, n = re.subn(rb"/Type\s*/Font", rb"/Type/Font 123 (x)", raw, count=1)
    assert n == 1
    f = tmp_path / "broken.pdf"
    f.write_bytes(raw)

    with caplog.at_level(logging.DEBUG, logger="pdf2video.parsers.pdf"):
        doc = load_document(f)
    assert "Readable body text survives the repair." in doc.chapters[0].text
    assert "MuPDF error" not in capfd.readouterr().err  # nothing printed by the C library
    assert any("invalid key in dict" in r.getMessage() for r in caplog.records)
    assert any(r.levelno == logging.INFO and "repaired" in r.getMessage() for r in caplog.records)


def test_pdf_vector_figure_with_margin_caption(tmp_path: Path):
    pdf = pymupdf.open()
    for title in ["First", "Second"]:
        page = pdf.new_page()
        page.insert_text((72, 90), title, fontsize=24)
        page.insert_text((150, 120), "Figure 9 shows how it works, as explained here.", fontsize=11)
    page = pdf[1]
    # A drawing with a label, and its caption in the left margin.
    shape = page.new_shape()
    shape.draw_rect(pymupdf.Rect(250, 300, 450, 450))
    shape.draw_line((250, 300), (450, 450))
    shape.draw_circle((350, 375), 40)
    shape.finish(color=(0, 0, 0), width=2)
    shape.commit()
    page.insert_text((260, 470), "Cost of the software", fontsize=10)
    page.insert_text((40, 308), "Figure 1.1", fontsize=9)
    page.insert_text((40, 330), "The cost grows faster.", fontsize=9)
    # A box elsewhere without a caption is not a figure.
    page.draw_rect(pymupdf.Rect(150, 150, 450, 200), color=(0, 0, 0))
    pdf.set_toc([[1, "First", 1], [1, "Second", 2]])
    f = tmp_path / "vector.pdf"
    pdf.save(f)
    doc = load_document(f)
    assert doc.chapters[0].images == []  # "Figure 9 shows …" is body text, not a caption
    [figure] = doc.chapters[1].images
    assert figure.caption == "Figure 1.1 The cost grows faster."
    crop = Image.open(io.BytesIO(figure.data))
    # The crop covers the drawing and its label (about 212×182 pt), not the page.
    assert 1.0 < crop.width / crop.height < 1.4 and max(crop.size) >= 800  # zoom ≤ 4
