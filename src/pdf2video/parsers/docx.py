"""DOCX loader based on python-docx; chapters come from Heading styles."""

from __future__ import annotations

import re
from pathlib import Path

import docx
from docx.oxml.ns import qn

from ..model import Chapter, ChapterImage, Document
from ..textutil import clean_paragraph

_HEADING = re.compile(r"^(heading|nadpis|überschrift|titre|título)\s*(\d)$", re.IGNORECASE)


def _heading_level(paragraph) -> int | None:
    name = (paragraph.style.name if paragraph.style is not None else "") or ""
    if name.lower() == "title":
        return 0
    if m := _HEADING.match(name.strip()):
        return int(m.group(2))
    outline = paragraph._p.find(f"{qn('w:pPr')}/{qn('w:outlineLvl')}")
    if outline is not None:
        return int(outline.get(qn("w:val"))) + 1
    return None


def load_docx(path: Path) -> Document:
    document = docx.Document(str(path))
    props = document.core_properties
    items: list[tuple[int | None, str, list[ChapterImage]]] = []
    for paragraph in document.paragraphs:
        text = clean_paragraph(paragraph.text)
        images: list[ChapterImage] = []
        for drawing in paragraph._p.xpath(".//w:drawing"):
            rids = drawing.xpath(".//a:blip/@r:embed")
            part = document.part.related_parts.get(rids[0]) if rids else None
            if part is not None and getattr(part, "blob", None):
                alt = drawing.xpath(".//wp:docPr/@descr")  # the picture's alt text
                images.append(ChapterImage(part.blob, clean_paragraph(alt[0] if alt else "")))
        if text or images:
            items.append((_heading_level(paragraph), text, images))

    title = (props.title or "").strip()
    if not title:
        title = next((t for lvl, t, _ in items if lvl == 0 and t), path.stem.replace("_", " "))

    levels = sorted({lvl for lvl, _, _ in items if lvl})
    chapter_level = next(
        (lvl for lvl in levels if sum(1 for x, *_ in items if x == lvl) >= 2),
        levels[0] if levels else None,
    )

    chapters: list[Chapter] = [Chapter(title, [])]
    for lvl, text, images in items:
        if lvl == 0:
            continue
        if chapter_level is not None and lvl is not None and lvl <= chapter_level and text:
            chapters.append(Chapter(text, []))
            continue
        if text:
            chapters[-1].paragraphs.append(text)
        chapters[-1].images.extend(images)
    return Document(path=path, title=title, author=(props.author or None), chapters=chapters)
